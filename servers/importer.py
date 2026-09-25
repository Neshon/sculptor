"""Импорт System BOM: состав изделия из Excel.

Файл устроен так: в первых строках — партномер изделия, дата и автор; затем
шапка таблицы; затем строки, разбитые на разделы (Boards, Cables,
Mechanics, Other) строками-разделителями из подчёркиваний. Раздел задаёт тип
позиции, колонка ``M/S`` отличает основную строку от замены — ровно как в
BOM платы.

Разбор отделён от записи: :func:`read` только читает файл и ничего не знает
про базу, :func:`apply_rows` переносит разобранное в позиции и строки
состава. Так разбор проверяется на файле без базы, а запись — на выдуманных
строках без файла.

Два правила, которые важнее прочего.

**Позиция заводится по номеру.** Если строка ссылается на то, чего в базе
ещё нет, заводится заготовка со статусом «в разработке»: карточку дополнят
позже, а состав не должен ждать. Опознавательный номер — OY P/N, а когда
его нет (вся механика от поставщика), номер поставщика.

**Повторный импорт заменяет только свои строки.** Строки, заведённые
руками, остаются на месте: иначе первая же повторная заливка стёрла бы
ручную работу. Отличаются они по полю ``source``.
"""

import re
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from components.matching import usable

from .models import BomLine, Item

SHEET = "System BOM"

# раздел файла -> тип позиции
KINDS = {
    "boards": Item.BOARD,
    "cables": Item.CABLE,
    "mechanics": Item.MECHANICAL,
    "mechanic": Item.MECHANICAL,
    "other": Item.OTHER,
    "materials": Item.MATERIAL,
}

# колонки шапки -> имена полей. Ищем по вхождению: в файлах встречаются
# «COMM/Vendor P/N:» с двоеточием и лишние пробелы
COLUMNS = (
    ("type", "type"),
    ("kind", "m/s"),
    ("gct_pn", "gct p/n"),
    ("oy_pn", "oy p/n"),
    ("vendor_pn", "vendor p/n"),
    ("vendor", "vendor"),
    ("description", "description"),
    ("comment", "comment"),
    ("quantity", "quantity"),
)

# «90 мм», «1,5 м», «3 шт»
QUANTITY = re.compile(r"^\s*(\d+(?:[.,]\d+)?)\s*([^\d\s]*)\s*$")


class ImportError_(Exception):
    """Файл не удалось разобрать — с объяснением для человека."""


def _text(value):
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def _content(value):
    """Значение ячейки, если в ней есть что-то кроме разделителя.

    Разделы в файле отбиты строкой из подчёркиваний, и подчёркивание стоит
    в каждой колонке — для проверки «есть ли в строке данные» такая ячейка
    должна считаться пустой.
    """
    text = usable(_text(value))
    return "" if not text.strip("_-—– ") else text


def _quantity(value):
    """Количество и единица измерения.

    В файле встречается и число, и запись вида «90 мм»: у ленты считают
    длину, а не штуки. Что не разобралось — единица, чтобы строка не
    потерялась; исходное значение остаётся в примечании.
    """
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)), "шт", ""

    text = _text(value)
    if not text:
        return Decimal("1"), "шт", ""

    match = QUANTITY.match(text)
    if not match:
        return Decimal("1"), "шт", text

    number, unit = match.groups()
    try:
        amount = Decimal(number.replace(",", "."))
    except InvalidOperation:
        return Decimal("1"), "шт", text
    return amount, (unit or "шт"), ""


def _header(rows):
    """Находит строку шапки и раскладку колонок."""
    for index, row in enumerate(rows):
        cells = [_text(cell).lower() for cell in row]
        if "m/s" in cells and any("p/n" in cell for cell in cells):
            layout = {}
            for name, marker in COLUMNS:
                for position, cell in enumerate(cells):
                    if marker in cell and name not in layout:
                        layout[name] = position
                        break
            return index, layout
    raise ImportError_("В файле нет строки шапки со столбцом «M/S».")


def read(source):
    """Разбирает файл. Возвращает (шапка изделия, строки)."""
    try:
        book = load_workbook(source, data_only=True, read_only=True)
    except Exception as exc:                       # noqa: BLE001
        raise ImportError_(f"Файл не читается: {exc}") from exc

    sheet = book[SHEET] if SHEET in book.sheetnames else book.worksheets[0]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    if not rows:
        raise ImportError_("Лист пуст.")

    start, layout = _header(rows)
    if "oy_pn" not in layout and "gct_pn" not in layout:
        raise ImportError_("В шапке нет ни OY PN, ни GCT PN.")

    # над шапкой стоят партномер изделия, дата и автор — по строке на каждое
    above = [_text(row[0]) for row in rows[:start] if any(row)]
    header = {
        "oy_pn": above[0] if above else "",
        "date": above[1] if len(above) > 1 else "",
        "author": above[2] if len(above) > 2 else "",
    }
    if not header["oy_pn"]:
        raise ImportError_("В первой строке нет OY PN позиции.")

    def cell(row, name):
        position = layout.get(name)
        if position is None or position >= len(row):
            return None
        return row[position]

    items, section = [], ""
    for offset, row in enumerate(rows[start + 1:], start=start + 2):
        section_cell = _text(cell(row, "type"))
        if section_cell and section_cell.strip("_ ") :
            section = section_cell

        kind = _text(cell(row, "kind")).upper()[:1]
        numbers = [_content(cell(row, name))
                   for name in ("oy_pn", "gct_pn", "vendor_pn")]
        described = _content(cell(row, "description"))

        if kind not in (BomLine.MAIN, BomLine.SUBSTITUTE):
            # Пометку M/S иногда забывают проставить — в разделе Other
            # такие строки встречаются. Раз номер и описание на месте,
            # это всё-таки строка состава: выбросить её значит потерять
            # позицию молча. Считаем основной, но помечаем в отчёте
            if not any(numbers) and not described:
                continue                  # разделители и пустые строки
            kind, assumed = BomLine.MAIN, True
        else:
            assumed = False

        quantity, unit, raw = _quantity(cell(row, "quantity"))
        items.append({
            "row": offset,
            "section": section,
            "kind": kind,
            "oy_pn": _content(cell(row, "oy_pn")),
            "gct_pn": _content(cell(row, "gct_pn")),
            "vendor_pn": _content(cell(row, "vendor_pn")),
            "vendor": _content(cell(row, "vendor")),
            "description": described,
            "comment": _content(cell(row, "comment")),
            "quantity": quantity,
            "unit": unit,
            "quantity_raw": raw,
            "kind_assumed": assumed,
        })

    if not items:
        raise ImportError_("В файле нет ни одной строки состава.")
    return header, items


def _find_item(oy_pn, gct_pn):
    """Позиция по номеру: сначала свой, потом поставщика."""
    if oy_pn:
        found = Item.objects.filter(oy_pn__iexact=oy_pn).first()
        if found:
            return found
    if gct_pn:
        found = Item.objects.filter(gct_pn__iexact=gct_pn).first()
        if found:
            return found
        # номер поставщика мог быть заведён как основной — так бывает у
        # механики, у которой своего номера нет вовсе
        return Item.objects.filter(oy_pn__iexact=gct_pn).first()
    return None


def apply_rows(header, rows, source, replace=True):
    """Переносит разобранное в базу. Возвращает отчёт."""
    parent = _find_item(header["oy_pn"], "")
    if parent is None:
        parent = Item.objects.create(
            oy_pn=header["oy_pn"], name=header["oy_pn"], kind=Item.SERVER,
            source=source)
        parent_created = True
    else:
        parent_created = False

    if replace:
        # только строки прошлых импортов: заведённые руками не трогаем
        BomLine.objects.filter(parent=parent).exclude(source="").delete()

    report = {"parent": parent, "parent_created": parent_created,
              "lines": 0, "items_created": [], "skipped": [], "assumed": []}

    for row in rows:
        number = row["oy_pn"] or row["gct_pn"]
        if not number and not row["description"]:
            report["skipped"].append((row["row"], "нет ни номера, ни описания"))
            continue

        child = _find_item(row["oy_pn"], row["gct_pn"]) if number else None
        if child is None and number:
            child = Item.objects.create(
                oy_pn=number,
                gct_pn=row["gct_pn"],
                name=row["description"][:255],
                kind=KINDS.get(row["section"].strip().lower(), Item.OTHER),
                status=Item.DRAFT,
                source=source)
            report["items_created"].append(child)

        if child is not None and child.pk == parent.pk:
            report["skipped"].append((row["row"], "строка ссылается на саму позицию"))
            continue

        if row["kind_assumed"]:
            report["assumed"].append(row["row"])

        comment = row["comment"]
        if row["quantity_raw"]:
            # количество не разобралось: сохраняем как было, чтобы человек
            # увидел исходную запись, а не подставленную единицу
            comment = f"{comment} (в файле количество: {row['quantity_raw']})".strip()

        BomLine.objects.create(
            parent=parent, child=child, kind=row["kind"],
            position=row["row"],
            oy_pn=row["oy_pn"], gct_pn=row["gct_pn"],
            description=row["description"],
            quantity=row["quantity"], unit=row["unit"],
            comment=comment, source=source, source_row=row["row"])
        report["lines"] += 1

    return report
