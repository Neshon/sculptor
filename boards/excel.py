"""Выгрузка состава платы в Excel.

Файл собирается в том же виде, в каком BOM приходят к нам: сверху шапка
парами «подпись — значение», ниже таблица состава. Подписи и заголовки
колонок взяты те же, что понимает ``boards.importer``, поэтому выгруженный
файл можно загрузить обратно — и получить ту же ревизию, а не мусор.
Круговой разбор проверен тестом ``RoundTripTests``.

Как и выгрузка CSV, Excel отдаёт **снимок файла**, а не подставленные из
библиотеки значения: это экспорт документа, и он должен совпадать с тем,
что импортировали.
"""

from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone

from components.export import EXPORT_LABELS, EXPORT_TIME_FORMAT
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHEET_NAME = "BOM"
CONTENT_TYPE = ("application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet")

# Шапка: подпись в файле → поле ревизии. Подписи должны оставаться теми,
# что перечислены в importer.HEADER_MAP, иначе обратный импорт их не узнает.
HEADER_ROWS = (
    ("Actual OY BOM P/N", "oy_pn"),
    ("Actual GCT BOM P/N", "gct_pcb"),
    ("Old OY BOM P/N", "previous_revision"),
    ("Decimal number", "decimal_pcba"),
)

# Подписи и формат времени отметки о выгрузке — общие с CSV
# (components/export.py): один и тот же состав приходит то в одном формате,
# то в другом, и называться отметка должна одинаково. Импорт эти строки не
# читает: в HEADER_MAP их нет, а шапка сопоставляется по подписям.

# Колонки таблицы: поле строки, заголовок, ширина, переносить ли текст.
# Заголовки — из importer.COLUMN_MAP, по той же причине.
COLUMNS = (
    ("position", "I/N", 6, False),
    ("kind", "M/S", 6, False),
    ("vendor_pn", "Vendor P/N", 26, False),
    ("vendor", "Vendor", 18, False),
    ("country", "Country", 14, False),
    ("oy_id", "OY ID", 16, False),
    ("oy_pn", "OY P/N", 20, False),
    ("gbt_pn", "GBT P/N", 20, False),
    ("group", "Group", 14, False),
    ("subgroup", "Subgroup", 16, False),
    ("description", "Description", 44, True),
    ("description_gbt", "Description GBT", 30, True),
    ("smt_tht", "SMT/THT", 10, False),
    ("references", "Part References", 30, True),
    ("qty", "QTY", 8, False),
    ("comment", "Comment", 24, True),
)

# цвета мини-брендбука OpenYard: тёмный фон шапки и вторичная подложка
DARK = "FF272B32"
LIGHT = "FFE4E4E4"
MUTED = "FFF0F0F0"
WHITE = "FFFFFFFF"

TITLE_FONT = Font(name="Verdana", size=12, bold=True, color=DARK)
LABEL_FONT = Font(name="Verdana", size=9, bold=True, color=DARK)
VALUE_FONT = Font(name="Verdana", size=9, color=DARK)
COLUMN_FONT = Font(name="Verdana", size=9, bold=True, color=WHITE)
CELL_FONT = Font(name="Verdana", size=9, color=DARK)

COLUMN_FILL = PatternFill("solid", fgColor=DARK)
LABEL_FILL = PatternFill("solid", fgColor=LIGHT)
SUBSTITUTE_FILL = PatternFill("solid", fgColor=MUTED)

THIN = Side(style="thin", color=LIGHT)
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TOP_LEFT = Alignment(horizontal="left", vertical="top")
TOP_LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="top")

# колонки, которые должны остаться числами, а не текстом
NUMERIC = {"position", "qty"}


def _cell_value(item, name):
    """Значение для ячейки: числа остаются числами, пустое — пустым.

    Что именно выгружается, решает сам BoardItem.export_value — там же
    записано, почему у замен пустой I/N.
    """
    value = item.export_value(name)
    if value is None or value == "":
        return None
    return value


def build_workbook(revision, items, exported_by="", exported_at=None):
    """Собирает книгу Excel с шапкой BOM и составом ревизии.

    ``exported_by`` и ``exported_at`` — кто и когда выгрузил файл. Без них
    отметка всё равно ставится, только с пустым именем: время известно
    всегда, а вот кто нажал кнопку — лишь когда выгрузку позвали из
    представления.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME

    row = _write_header(sheet, revision, exported_by, exported_at)
    row = _write_columns(sheet, row + 1)
    _write_items(sheet, row, items)
    _finish(sheet, row - 1, len(items))
    return workbook


def _write_header(sheet, revision, exported_by="", exported_at=None):
    """Название и пары «подпись — значение». Возвращает номер последней строки."""
    title = sheet.cell(row=1, column=1,
                       value=f"Спецификация платы {revision.oy_pn}".strip())
    title.font = TITLE_FONT
    title.alignment = TOP_LEFT

    row = 2
    for label, name in HEADER_ROWS:
        row += 1
        cell = sheet.cell(row=row, column=1, value=label)
        cell.font = LABEL_FONT
        cell.fill = LABEL_FILL
        cell.alignment = TOP_LEFT
        # значение обязательно правее подписи: импорт ищет его именно там
        value = sheet.cell(row=row, column=2, value=getattr(revision, name, "") or "")
        value.font = VALUE_FONT
        value.alignment = TOP_LEFT

    when = exported_at or timezone.localtime()
    for label, text in zip(EXPORT_LABELS,
                           (when.strftime(EXPORT_TIME_FORMAT), exported_by)):
        row += 1
        cell = sheet.cell(row=row, column=1, value=label)
        cell.font = LABEL_FONT
        cell.fill = LABEL_FILL
        cell.alignment = TOP_LEFT
        value = sheet.cell(row=row, column=2, value=text or "")
        value.font = VALUE_FONT
        value.alignment = TOP_LEFT
    return row


def _write_columns(sheet, row):
    """Заголовки таблицы. Возвращает номер первой строки состава."""
    for index, (_, label, width, _wrap) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=row, column=index, value=label)
        cell.font = COLUMN_FONT
        cell.fill = COLUMN_FILL
        cell.alignment = CENTER
        sheet.column_dimensions[get_column_letter(index)].width = width
    return row + 1


def _write_items(sheet, first_row, items):
    for offset, item in enumerate(items):
        row = first_row + offset
        substitute = not item.is_main
        for index, (name, _label, _width, wrap) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row, column=index,
                              value=_cell_value(item, name))
            cell.font = CELL_FONT
            cell.border = CELL_BORDER
            cell.alignment = TOP_LEFT_WRAP if wrap else (
                CENTER if name in NUMERIC or name == "kind" else TOP_LEFT)
            if substitute:
                # замены подсвечены так же, как в таблице на экране
                cell.fill = SUBSTITUTE_FILL


def _finish(sheet, columns_row, count):
    """Закрепление шапки и фильтр по колонкам — с этим файлом ещё работают."""
    sheet.freeze_panes = sheet.cell(row=columns_row + 1, column=1)
    last_column = get_column_letter(len(COLUMNS))
    last_row = columns_row + count
    sheet.auto_filter.ref = f"A{columns_row}:{last_column}{max(last_row, columns_row)}"


def workbook_response(filename, workbook):
    """Отдаёт книгу на скачивание, не сохраняя её на диск."""
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type=CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
